# -*- coding: utf-8 -*-
"""Office 文件转 Markdown：DOCX / DOC / XLSX / XLS。

- DOCX：mammoth 转 HTML，再转 Markdown；表格保留 HTML 合并单元格。
- XLSX：openpyxl 按 sheet 转 Markdown 管道表。
- DOC/XLS：先用 LibreOffice headless 转成 DOCX/XLSX，再走上面的流程。
"""
from __future__ import annotations

import html as htmlmod
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from datetime import date, datetime
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag

CREATE_NO_WINDOW = 0x08000000 if os.name == "nt" else 0
SUPPORTED = {".pdf", ".docx", ".doc", ".xlsx", ".xls"}
OFFICE = {".docx", ".doc", ".xlsx", ".xls"}


def _safe(s):
    if not isinstance(s, str):
        return s
    return s.encode("utf-16-le", "surrogatepass").decode("utf-16-le", "replace")


def file_format(path) -> str:
    return Path(path).suffix.lower()


def is_supported(path) -> bool:
    return file_format(path) in SUPPORTED


def find_soffice() -> Path | None:
    cands = []
    env = os.environ.get("SOFFICE")
    if env:
        cands.append(Path(env))
    which = shutil.which("soffice")
    if which:
        cands.append(Path(which))
    cands += [
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ]
    for p in cands:
        try:
            if p and p.is_file():
                return p
        except Exception:
            pass
    return None


# ---------------- 通用 ----------------
def _out_paths(path: Path, out_dir: Path, out_stem: str | None):
    stem = out_stem or path.stem
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / (stem + ".md"), out_dir / (stem + "_qa.json")


def _count_tables(md: str) -> int:
    return md.count("<table") + len(re.findall(r"^\| *---", md, re.M))


def _write_result(path: Path, out_dir: Path, out_stem: str | None, md: str,
                  qa: dict, on_event):
    md_path, qa_path = _out_paths(path, out_dir, out_stem)
    on_event("phase", {"name": "write"})
    md_path.write_text(_safe(md), encoding="utf-8")
    qa_path.write_text(_safe(json.dumps(qa, ensure_ascii=False, indent=2)),
                       encoding="utf-8")
    on_event("done", {"path": str(md_path), "qa": qa})
    return {"path": str(md_path), "qa": qa}


# ---------------- DOCX ----------------
def _inline(node) -> str:
    out = []
    for child in node.children:
        if isinstance(child, NavigableString):
            out.append(str(child))
        elif isinstance(child, Tag):
            t = child.name.lower()
            if t in ("strong", "b"):
                out.append(f"**{_inline(child).strip()}**")
            elif t in ("em", "i"):
                out.append(f"*{_inline(child).strip()}*")
            elif t == "code":
                out.append(f"`{_inline(child).strip()}`")
            elif t == "a":
                href = child.get("href", "")
                out.append(f"[{_inline(child).strip()}]({href})")
            elif t == "br":
                out.append("\n")
            elif t == "img":
                alt = child.get("alt", "").strip()
                if alt:
                    out.append(f"![{alt}]()")
            else:
                out.append(_inline(child))
    return "".join(out)


def _html_to_md(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    body = soup.body or soup
    blocks = []

    def walk(container):
        for node in container.children:
            if isinstance(node, NavigableString):
                text = str(node).strip()
                if text:
                    blocks.append(text)
                continue
            if not isinstance(node, Tag):
                continue
            t = node.name.lower()
            if re.fullmatch(r"h[1-6]", t):
                level = int(t[1])
                txt = _inline(node).strip()
                if txt:
                    blocks.append("#" * level + " " + txt)
            elif t == "p":
                txt = _inline(node).strip()
                if txt:
                    blocks.append(txt)
            elif t == "table":
                blocks.append(str(node))
            elif t in ("ul", "ol"):
                items = []
                for i, li in enumerate(node.find_all("li", recursive=False), 1):
                    txt = _inline(li).strip()
                    if txt:
                        prefix = "- " if t == "ul" else f"{i}. "
                        items.append(prefix + txt)
                if items:
                    blocks.append("\n".join(items))
            elif t == "blockquote":
                txt = _inline(node).strip()
                if txt:
                    blocks.append("\n".join("> " + x for x in txt.splitlines()))
            elif t == "pre":
                blocks.append("```\n" + node.get_text("\n", strip=False) + "\n```")
            elif t == "img":
                alt = node.get("alt", "").strip()
                if alt:
                    blocks.append(f"![{alt}]()")
            else:
                walk(node)

    walk(body)
    out = "\n\n".join(blocks)
    return re.sub(r"\n{3,}", "\n\n", out).strip() + "\n"


def _docx_fallback(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            txt = p.text.strip()
            if not txt:
                continue
            style = (p.style.name if p.style else "").lower()
            if style.startswith("heading"):
                m = re.search(r"(\d+)", style)
                level = min(6, int(m.group(1))) if m else 1
                blocks.append("#" * level + " " + txt)
            elif "list" in style:
                blocks.append("- " + txt)
            else:
                blocks.append(txt)
        elif tag == "tbl":
            tb = Table(child, doc)
            rows = []
            for row in tb.rows:
                cells = []
                for cell in row.cells:
                    v = cell.text.replace("\n", "<br>").replace("|", "\\|").strip()
                    cells.append(v)
                rows.append(cells)
            if rows:
                n = max(len(r) for r in rows)
                rows = [r + [""] * (n - len(r)) for r in rows]
                lines = ["| " + " | ".join(rows[0]) + " |",
                         "| " + " | ".join(["---"] * n) + " |"]
                lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                blocks.append("\n".join(lines))
    return "\n\n".join(blocks).strip() + "\n"


def docx_to_md(path: Path) -> tuple[str, list[str]]:
    warnings = []
    try:
        import mammoth
        with open(path, "rb") as fh:
            result = mammoth.convert_to_html(
                fh, style_map=("p[style-name='Title'] => h1:fresh" + chr(10) +
                              "p[style-name='Subtitle'] => p:fresh"))
        warnings = [str(m).strip() for m in (result.messages or []) if str(m).strip()]
        return _html_to_md(result.value), warnings
    except Exception as e:
        warnings.append(f"mammoth 失败，改用 python-docx：{e}")
        return _docx_fallback(path), warnings


def convert_docx(path: Path, out_dir: Path, out_stem: str | None = None,
                 on_event=None, source_path: Path | None = None) -> dict:
    ev = on_event or (lambda *a, **k: None)
    t0 = time.time()
    ev("phase", {"name": "office_parse", "format": "docx"})
    md, warnings = docx_to_md(path)
    source = source_path or path
    header = f"<!-- source: {source.name} -->\n\n"
    md = header + md
    qa = {
        "pages": 0, "parsed_pages": 0, "text_pages": 0,
        "chars": len(md), "tables_mid": 0,
        "tables_md": _count_tables(md),
        "min_page_recall": None, "md_pdf_chars_ratio": None,
        "avg_page_recall": None, "pages_below_0.9": 0,
        "suspicious_cells": 0, "pua_left": 0,
        "source": str(source), "format": "docx",
        "warnings": warnings[:20],
        "total_sec": round(time.time() - t0, 1),
    }
    return _write_result(source, out_dir, out_stem, md, qa, ev)


# ---------------- XLSX ----------------
def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat(sep=" ")
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.15g}"
    s = str(v)
    return s.replace("\n", "<br>").replace("|", "\\|").strip()


def xlsx_to_md(path: Path) -> tuple[str, dict]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=False)
    sheets = len(wb.sheetnames)
    parts = []
    n_tables = 0
    for ws in wb.worksheets:
        parts.append(f"## {ws.title}")
        max_r, max_c = ws.max_row, ws.max_column
        if not max_r or not max_c:
            parts.append("（空工作表）")
            continue
        merged = {}
        for rng in ws.merged_cells.ranges:
            try:
                v = ws.cell(rng.min_row, rng.min_col).value
            except Exception:
                v = None
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged[(r, c)] = v
        rows = []
        for r in range(1, max_r + 1):
            row = []
            for c in range(1, max_c + 1):
                v = merged.get((r, c), ws.cell(r, c).value)
                row.append(_cell_text(v))
            if any(x for x in row):
                rows.append(row)
        if not rows:
            parts.append("（空工作表）")
            continue
        n = max(len(r) for r in rows)
        rows = [r + [""] * (n - len(r)) for r in rows]
        header = rows[0]
        # 表头空列补“列N”
        header = [h if h else f"列{i+1}" for i, h in enumerate(header)]
        lines = ["| " + " | ".join(header) + " |",
                 "| " + " | ".join(["---"] * n) + " |"]
        lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        parts.append("\n".join(lines))
        n_tables += 1
    wb.close()
    return "\n\n".join(parts).strip() + "\n", {"sheets": len(wb.sheetnames),
                                               "tables": n_tables}


def convert_xlsx(path: Path, out_dir: Path, out_stem: str | None = None,
                 on_event=None, source_path: Path | None = None) -> dict:
    ev = on_event or (lambda *a, **k: None)
    t0 = time.time()
    ev("phase", {"name": "office_parse", "format": "xlsx"})
    body, info = xlsx_to_md(path)
    source = source_path or path
    md = f"<!-- source: {source.name} -->\n\n" + body
    qa = {
        "pages": 0, "parsed_pages": 0, "text_pages": 0,
        "chars": len(md), "tables_mid": info["tables"],
        "tables_md": _count_tables(md),
        "min_page_recall": None, "md_pdf_chars_ratio": None,
        "avg_page_recall": None, "pages_below_0.9": 0,
        "suspicious_cells": 0, "pua_left": 0,
        "source": str(source), "format": "xlsx",
        "sheets": info["sheets"],
        "total_sec": round(time.time() - t0, 1),
    }
    return _write_result(source, out_dir, out_stem, md, qa, ev)


# ---------------- 旧版 DOC / XLS ----------------
def _legacy_to_modern(path: Path, tmp_dir: Path, on_event) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(
            "旧版 DOC/XLS 需要 LibreOffice。请在主界面点“安装/修复”，"
            "或在菜单里安装 LibreOffice 后重试。")
    target = "docx" if path.suffix.lower() == ".doc" else "xlsx"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    profile = tmp_dir / "lo_profile"
    profile.mkdir(parents=True, exist_ok=True)
    on_event("phase", {"name": "office_convert", "format": path.suffix.lower()})
    cmd = [str(soffice), "--headless",
           f"-env:UserInstallation={profile.as_uri()}",
           "--convert-to", target, "--outdir", str(tmp_dir), str(path)]
    p = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                       timeout=600, creationflags=CREATE_NO_WINDOW)
    if p.returncode != 0:
        tail = (p.stdout or b"").decode("utf-8", "replace")[-800:]
        raise RuntimeError(f"LibreOffice 转换失败 rc={p.returncode}: {tail}")
    wanted = tmp_dir / (path.stem + "." + target)
    if wanted.is_file():
        return wanted
    found = sorted(tmp_dir.glob("*." + target))
    if found:
        return found[0]
    raise RuntimeError("LibreOffice 没有生成目标文件，请检查文件是否损坏")


def convert_legacy(path: Path, out_dir: Path, out_stem: str | None = None,
                   tmp_dir: Path | None = None, on_event=None) -> dict:
    ev = on_event or (lambda *a, **k: None)
    work = Path(tmp_dir or (out_dir / "_legacy_tmp"))
    work.mkdir(parents=True, exist_ok=True)
    modern = _legacy_to_modern(path, work, ev)
    try:
        if modern.suffix.lower() == ".docx":
            return convert_docx(modern, out_dir, out_stem=out_stem,
                                on_event=ev, source_path=path)
        return convert_xlsx(modern, out_dir, out_stem=out_stem,
                            on_event=ev, source_path=path)
    finally:
        shutil.rmtree(work, ignore_errors=True)


# ---------------- 统一入口 ----------------
def convert_any(path, out_dir, overwrite=False, keep_json=False, pages="all",
                device="auto", tier="basic", tmp_dir=None, on_event=None,
                out_stem=None, page_workers=1, cpu_threads=0):
    path = Path(path)
    out_dir = Path(out_dir)
    fmt = file_format(path)
    if fmt == ".pdf":
        from convert import convert_one
        return convert_one(path, out_dir, overwrite=overwrite,
                           keep_json=keep_json, pages=pages, device=device,
                           tier=tier, tmp_dir=tmp_dir, on_event=on_event,
                           out_stem=out_stem, page_workers=page_workers,
                           cpu_threads=cpu_threads)
    md_path, _ = _out_paths(path, out_dir, out_stem)
    if md_path.exists() and not overwrite:
        if on_event:
            on_event("skipped", {"path": str(md_path)})
        return {"path": str(md_path), "skipped": True}
    if fmt == ".docx":
        return convert_docx(path, out_dir, out_stem=out_stem,
                            on_event=on_event)
    if fmt == ".xlsx":
        return convert_xlsx(path, out_dir, out_stem=out_stem,
                            on_event=on_event)
    if fmt in (".doc", ".xls"):
        return convert_legacy(path, out_dir, out_stem=out_stem,
                              tmp_dir=Path(tmp_dir) if tmp_dir else None,
                              on_event=on_event)
    raise ValueError(f"不支持的文件类型：{fmt or '未知'}")
